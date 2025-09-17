#!/usr/bin/env python3
"""
World-Class DOCX to HTML Converter for Sharedo Templates
Optimized for email compatibility with template variable preservation
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from bs4 import BeautifulSoup
import html

class ShareDoDocxToHtmlConverter:
    """Advanced DOCX to HTML converter optimized for Sharedo email templates"""
    
    def __init__(self):
        self.template_patterns = {
            'variable': re.compile(r'\[_+\]'),  # Matches [_____]
            'if_start': re.compile(r'\{\{#if\s+(.*?)\}\}', re.IGNORECASE),
            'if_end': re.compile(r'\{\{/if\}\}', re.IGNORECASE),
            'each_start': re.compile(r'\{\{#each\s+(.*?)\}\}', re.IGNORECASE),
            'each_end': re.compile(r'\{\{/each\}\}', re.IGNORECASE),
            'handlebars': re.compile(r'\{\{(.*?)\}\}'),
            'merge_field': re.compile(r'\[\[(.*?)\]\]'),
        }
        
        self.email_safe_styles = {
            'font-family': 'Arial, Helvetica, sans-serif',
            'color': '#333333',
            'line-height': '1.6',
        }
    
    def convert(self, docx_path, output_path=None):
        """Main conversion method"""
        doc = Document(docx_path)
        html_content = self._generate_html(doc)
        
        if output_path:
            Path(output_path).write_text(html_content, encoding='utf-8')
            print(f"✅ HTML file saved to: {output_path}")
        
        return html_content
    
    def _generate_html(self, doc):
        """Generate complete HTML document optimized for email"""
        html_parts = []
        
        # Email-optimized HTML header
        html_parts.append(self._get_html_header())
        
        # Process document content
        body_content = self._process_document_body(doc)
        html_parts.append(body_content)
        
        # HTML footer
        html_parts.append(self._get_html_footer())
        
        # Beautify and optimize
        soup = BeautifulSoup(''.join(html_parts), 'html.parser')
        return soup.prettify()
    
    def _get_html_header(self):
        """Email-optimized HTML header with responsive meta tags"""
        return '''<!DOCTYPE html>
<html lang="en" xmlns="http://www.w3.org/1999/xhtml" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <meta http-equiv="X-UA-Compatible" content="IE=edge">
    <meta name="x-apple-disable-message-reformatting">
    <title>Sharedo Template</title>
    <!--[if mso]>
    <noscript>
        <xml>
            <o:OfficeDocumentSettings>
                <o:AllowPNG/>
                <o:PixelsPerInch>96</o:PixelsPerInch>
            </o:OfficeDocumentSettings>
        </xml>
    </noscript>
    <![endif]-->
    <style type="text/css">
        /* Email Client Reset Styles */
        body, table, td, a { -webkit-text-size-adjust: 100%; -ms-text-size-adjust: 100%; }
        table, td { mso-table-lspace: 0pt; mso-table-rspace: 0pt; }
        img { -ms-interpolation-mode: bicubic; border: 0; outline: none; text-decoration: none; }
        
        /* Remove default styling */
        body { margin: 0 !important; padding: 0 !important; width: 100% !important; min-width: 100% !important; }
        
        /* Base styles */
        body {
            font-family: Arial, Helvetica, sans-serif;
            font-size: 16px;
            line-height: 1.6;
            color: #333333;
            background-color: #f4f4f4;
        }
        
        /* Container styles - matching Word document dimensions */
        .email-container {
            max-width: 816px; /* 8.5 inches at 96 DPI */
            margin: 0 auto;
            background-color: #ffffff;
            padding: 96px; /* 1 inch margins like Word */
            box-shadow: 0 0 10px rgba(0,0,0,0.1);
        }
        
        /* Inner content width matches Word's 6.5 inch text width */
        .email-content {
            max-width: 624px; /* 6.5 inches at 96 DPI */
            margin: 0 auto;
        }
        
        /* Typography */
        h1, h2, h3, h4, h5, h6 {
            margin: 0 0 10px 0;
            font-weight: bold;
            line-height: 1.2;
        }
        
        p {
            margin: 0 0 15px 0;
            padding: 0;
        }
        
        /* Template variables */
        .template-variable {
            background-color: #fffbdd;
            padding: 2px 4px;
            border-radius: 3px;
            font-family: 'Courier New', monospace;
            color: #d73a49;
        }
        
        .template-if {
            border-left: 3px solid #0969da;
            padding-left: 10px;
            margin: 10px 0;
        }
        
        /* Tables */
        .data-table {
            width: 100%;
            border-collapse: collapse;
            margin: 15px 0;
        }
        
        .data-table th,
        .data-table td {
            padding: 10px;
            border: 1px solid #dddddd;
            text-align: left;
        }
        
        .data-table th {
            background-color: #f8f9fa;
            font-weight: bold;
        }
        
        /* Mobile Responsive */
        @media screen and (max-width: 840px) {
            .email-container {
                width: 100% !important;
                padding: 20px !important;
                box-shadow: none !important;
            }
            
            table {
                width: 100% !important;
            }
            
            td {
                padding: 10px !important;
            }
        }
    </style>
</head>
<body style="margin: 0; padding: 0; background-color: #f4f4f4;">
    <!--[if mso]>
    <table role="presentation" border="0" cellpadding="0" cellspacing="0" width="100%">
    <tr>
    <td>
    <![endif]-->
    
    <div class="email-container" style="max-width: 816px; margin: 0 auto; background-color: #ffffff; padding: 96px; box-shadow: 0 0 10px rgba(0,0,0,0.1);">
        <div class="email-content" style="max-width: 624px; margin: 0 auto;">
        <table role="presentation" border="0" cellpadding="0" cellspacing="0" width="100%">
            <tr>
                <td style="padding: 20px;">'''
    
    def _get_html_footer(self):
        """HTML footer"""
        return '''
                </td>
            </tr>
        </table>
        </div>
    </div>
    
    <!--[if mso]>
    </td>
    </tr>
    </table>
    <![endif]-->
</body>
</html>'''
    
    def _process_document_body(self, doc):
        """Process document paragraphs and tables"""
        content_parts = []
        
        for element in self._iter_block_items(doc):
            if hasattr(element, 'text'):  # Paragraph
                para_html = self._process_paragraph(element)
                if para_html:
                    content_parts.append(para_html)
            elif hasattr(element, 'rows'):  # Table
                table_html = self._process_table(element)
                if table_html:
                    content_parts.append(table_html)
        
        return '\n'.join(content_parts)
    
    def _iter_block_items(self, document):
        """Yield each paragraph and table in document order"""
        from docx.document import Document as DocumentType
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        
        parent = document if isinstance(document, DocumentType) else document._element
        
        if hasattr(parent, 'element'):
            parent_elm = parent.element.body
        else:
            parent_elm = parent.body if hasattr(parent, 'body') else parent
        
        for child in parent_elm.iterchildren():
            if child.tag.endswith('p'):
                yield Paragraph(child, document if isinstance(document, DocumentType) else document._parent)
            elif child.tag.endswith('tbl'):
                yield Table(child, document if isinstance(document, DocumentType) else document._parent)
    
    def _process_paragraph(self, paragraph):
        """Process a single paragraph with formatting and template preservation"""
        if not paragraph.text.strip():
            return ''
        
        # Detect paragraph alignment
        alignment = self._get_alignment_style(paragraph)
        
        # Process runs (text with formatting)
        content_parts = []
        for run in paragraph.runs:
            run_html = self._process_run(run)
            if run_html:
                content_parts.append(run_html)
        
        if not content_parts:
            return ''
        
        content = ''.join(content_parts)
        
        # Detect and wrap special content
        if self._is_heading(paragraph):
            level = self._get_heading_level(paragraph)
            return f'<h{level} style="{alignment}margin: 20px 0 10px 0;">{content}</h{level}>'
        else:
            return f'<p style="{alignment}margin: 0 0 15px 0; line-height: 1.6;">{content}</p>'
    
    def _process_run(self, run):
        """Process a text run with formatting"""
        if not run.text:
            return ''
        
        text = run.text
        
        # Preserve template variables
        text = self._preserve_template_variables(text)
        
        # Apply formatting
        styles = []
        
        if run.bold:
            text = f'<strong>{text}</strong>'
        if run.italic:
            text = f'<em>{text}</em>'
        if run.underline:
            text = f'<u>{text}</u>'
        
        # Font size
        if run.font.size:
            size_pt = run.font.size.pt if hasattr(run.font.size, 'pt') else 12
            styles.append(f'font-size: {size_pt}pt')
        
        # Font color
        if run.font.color and run.font.color.rgb:
            color_hex = self._rgb_to_hex(run.font.color.rgb)
            styles.append(f'color: {color_hex}')
        
        # Font family
        if run.font.name:
            styles.append(f'font-family: {run.font.name}, Arial, sans-serif')
        
        if styles:
            style_str = '; '.join(styles)
            text = f'<span style="{style_str}">{text}</span>'
        
        return text
    
    def _preserve_template_variables(self, text):
        """Preserve and highlight template variables"""
        # Preserve [_____] placeholders
        text = self.template_patterns['variable'].sub(
            lambda m: f'<span class="template-variable" style="background-color: #fffbdd; padding: 2px 4px; border-radius: 3px; font-family: monospace; color: #d73a49;">{m.group()}</span>',
            text
        )
        
        # Preserve {{variables}}
        text = self.template_patterns['handlebars'].sub(
            lambda m: f'<span class="template-variable" style="background-color: #e6f7ff; padding: 2px 4px; border-radius: 3px; font-family: monospace; color: #0969da;">{html.escape(m.group())}</span>',
            text
        )
        
        # Preserve [[merge fields]]
        text = self.template_patterns['merge_field'].sub(
            lambda m: f'<span class="template-variable" style="background-color: #fff0f0; padding: 2px 4px; border-radius: 3px; font-family: monospace; color: #cf222e;">{m.group()}</span>',
            text
        )
        
        return text
    
    def _process_table(self, table):
        """Convert table to HTML with email-safe styling"""
        html_parts = ['<table class="data-table" style="width: 100%; border-collapse: collapse; margin: 15px 0;" cellpadding="0" cellspacing="0">']
        
        for row_idx, row in enumerate(table.rows):
            html_parts.append('<tr>')
            for cell in row.cells:
                tag = 'th' if row_idx == 0 else 'td'
                style = 'padding: 10px; border: 1px solid #dddddd; text-align: left;'
                if row_idx == 0:
                    style += ' background-color: #f8f9fa; font-weight: bold;'
                
                cell_content = ' '.join([para.text for para in cell.paragraphs])
                cell_content = self._preserve_template_variables(cell_content)
                
                html_parts.append(f'<{tag} style="{style}">{cell_content}</{tag}>')
            html_parts.append('</tr>')
        
        html_parts.append('</table>')
        return '\n'.join(html_parts)
    
    def _get_alignment_style(self, paragraph):
        """Get CSS alignment from paragraph alignment"""
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            return 'text-align: center; '
        elif paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            return 'text-align: right; '
        elif paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
            return 'text-align: justify; '
        return 'text-align: left; '
    
    def _is_heading(self, paragraph):
        """Check if paragraph is a heading"""
        return paragraph.style.name.startswith('Heading') if paragraph.style else False
    
    def _get_heading_level(self, paragraph):
        """Get heading level from style name"""
        if paragraph.style and paragraph.style.name.startswith('Heading'):
            try:
                return int(paragraph.style.name.replace('Heading ', ''))
            except:
                return 2
        return 2
    
    def _rgb_to_hex(self, rgb):
        """Convert RGBColor to hex string"""
        if isinstance(rgb, RGBColor):
            return '#{:02x}{:02x}{:02x}'.format(
                rgb[0] if rgb[0] else 0,
                rgb[1] if rgb[1] else 0,
                rgb[2] if rgb[2] else 0
            )
        return '#000000'


def main():
    """Main conversion function"""
    converter = ShareDoDocxToHtmlConverter()
    
    # Input and output paths
    docx_file = "We refer to the telephone conversation.docx"
    html_file = "Output/We_refer_to_telephone_conversation.html"
    
    print("🚀 Starting DOCX to HTML Conversion")
    print("=" * 50)
    
    # Perform conversion
    html_content = converter.convert(docx_file, html_file)
    
    print("\n📊 Conversion Statistics:")
    print(f"  • Input: {docx_file}")
    print(f"  • Output: {html_file}")
    print(f"  • HTML Size: {len(html_content)} characters")
    
    # Count preserved template variables
    variable_count = len(re.findall(r'\[_+\]', html_content))
    print(f"  • Template Variables Preserved: {variable_count}")
    
    print("\n✅ Conversion completed successfully!")
    print("📧 HTML is optimized for email clients including Outlook")
    
    return html_content


if __name__ == "__main__":
    main()